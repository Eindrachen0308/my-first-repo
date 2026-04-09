from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

doc = Document()

# -- Page margins --
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

# -- Style configuration --
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = RGBColor(0x33, 0x33, 0x33)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for level in range(1, 4):
    h_style = doc.styles[f'Heading {level}']
    h_style.font.name = 'Calibri'
    h_style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

doc.styles['Heading 1'].font.size = Pt(22)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(24)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)

doc.styles['Heading 2'].font.size = Pt(16)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(18)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(8)

doc.styles['Heading 3'].font.size = Pt(13)
doc.styles['Heading 3'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 3'].paragraph_format.space_after = Pt(6)

# Configure List Bullet style
list_style = doc.styles['List Bullet']
list_style.font.name = 'Calibri'
list_style.font.size = Pt(11)
list_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
list_style.paragraph_format.space_after = Pt(3)

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.bold = True
        run_b.font.name = 'Calibri'
        run_b.font.size = Pt(11)
        run_n = p.add_run(text)
        run_n.font.name = 'Calibri'
        run_n.font.size = Pt(11)
    else:
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
    return p

def add_separator(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run('─' * 60)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    run.font.size = Pt(8)

# ============================================================
# JD 1: Software Engineer, Backend
# ============================================================

doc.add_heading('Software Engineer, Backend', level=1)

doc.add_heading('Mission', level=2)
doc.add_paragraph(
    'Design, build, and operate the backend infrastructure that powers Shizuku\'s conversational AI, '
    'live-streaming capabilities, and social media integrations. Architect a robust orchestration layer '
    'that seamlessly unifies LLM, TTS, ASR, and other AI services into a cohesive product experience. '
    'Deliver sub-300ms response times at 500+ concurrent YouTube connections, enabling real-time '
    'interactions that feel natural and immediate.'
)

doc.add_heading('About Shizuku', level=2)
doc.add_paragraph(
    'Shizuku is a Japan-born AI companion active across YouTube and X (formerly Twitter). Already '
    'operating live streams and cultivating an engaged community, Shizuku is now entering its next '
    'growth phase. Backed by a16z in their first-ever investment in a Japanese startup at the seed '
    'stage, we are on a mission to bring Japan\'s unique fusion of entertainment and AI to a global audience.'
)

doc.add_heading('Team Structure', level=2)
doc.add_paragraph(
    'You will join a founding-stage team led by Aki, co-founder, ML engineer, and researcher '
    '(formerly at Meta and Luma AI), alongside our Engineering Director. As the first dedicated backend '
    'engineer, you will have a direct hand in technology selection, architectural decisions, and product '
    'direction\u2014collaborating closely with the founding members on every major decision.'
)

doc.add_heading('Key Responsibilities', level=2)
bullets = [
    'Design, implement, and operate backend services and APIs that underpin Shizuku\'s conversation engine, live-streaming pipeline, and social media integrations',
    'Build real-time integration infrastructure for external platforms including YouTube, X, and Discord',
    'Architect and implement an orchestration layer that integrates LLM, TTS, ASR, and other AI capabilities into the product',
    'Design low-latency architectures capable of sustaining 500+ concurrent connections with sub-300ms response times',
    'Architect, provision, and manage production environments on AWS',
    'Establish and continuously improve asynchronous job processing, state management, observability and alerting, testing strategies, and CI/CD pipelines',
]
for b in bullets:
    add_bullet(doc, b)

doc.add_heading('Required Qualifications', level=2)
bullets = [
    'Approximately 3+ years of backend development experience across multiple languages, including Python or TypeScript',
    'Proven experience designing and operating systems involving Web APIs, asynchronous processing, and persistent data stores',
    'Demonstrated expertise in designing and implementing systems with stringent real-time and low-latency requirements',
    'Track record of translating ambiguous requirements into well-defined specifications and independently driving projects from design through release',
    'Ability to work on-site at our Tokyo office (primarily in-office with flexible remote arrangements)',
]
for b in bullets:
    add_bullet(doc, b)

doc.add_heading('Preferred Qualifications', level=2)
bullets = [
    'Hands-on experience with AWS infrastructure provisioning and managed services',
    'Experience developing AI/LLM-powered products or integrating AI services via APIs',
    'Experience with LLM post-training, evaluation, or building training pipelines using frameworks such as ms-swift or Megatron-LM',
    'Experience building training data curation and evaluation infrastructure',
    'Implementation experience with external platform APIs (YouTube, X, Discord, etc.)',
    'Experience improving operational excellence through log architecture, monitoring, and SLI/SLO management',
    'Experience operating consumer-facing (toC) products at scale',
    'Experience with 0-to-1 product development in a startup or small-team environment',
    'Ability to read and write technical documentation in English (the team currently operates primarily in Japanese, with plans to transition to a global working environment in the medium term)',
]
for b in bullets:
    add_bullet(doc, b)

doc.add_heading('Who You Are', level=2)
bullets = [
    ('Learning Agility \u2014 ', 'You embrace unfamiliar domains with enthusiasm, rapidly acquiring new knowledge and translating it into tangible outcomes.'),
    ('Purpose-Driven Ownership \u2014 ', 'You work backward from user value and business impact, proactively identifying challenges and driving them to completion.'),
    ('Comfort with Ambiguity \u2014 ', 'You thrive in environments without a dedicated PM, engaging directly with founding members to shape product decisions and move forward decisively.'),
    ('Humility & Respect \u2014 ', 'You value diverse perspectives and communicate with integrity and constructiveness.'),
    ('Resilience \u2014 ', 'You remain tenacious and optimistic in the face of adversity, consistently finding a way through.'),
]
for prefix, text in bullets:
    add_bullet(doc, text, bold_prefix=prefix)

# Page break
doc.add_page_break()

# ============================================================
# JD 2: ML Engineer
# ============================================================

doc.add_heading('ML Engineer', level=1)

doc.add_heading('Mission', level=2)
doc.add_paragraph(
    'Lead the research and development of the AI models that define Shizuku\'s voice and intelligence. '
    'With a primary focus on advancing our Text-to-Speech (TTS) technology, you will also drive '
    'multimodal expansion across NLP, speech recognition, and\u2014looking ahead\u2014computer vision and '
    'humanoid robotics, continuously elevating Shizuku\'s expressive capabilities. Balance the improvement '
    'of production TTS models with the exploration and development of next-generation architectures, '
    'while owning the MLOps lifecycle to ensure continuous, measurable progress.'
)

doc.add_heading('About Shizuku', level=2)
doc.add_paragraph(
    'Shizuku is a Japan-born AI companion active across YouTube and X (formerly Twitter). Already '
    'operating live streams and cultivating an engaged community, Shizuku is now entering its next '
    'growth phase. Backed by a16z in their first-ever investment in a Japanese startup at the seed '
    'stage, we are on a mission to bring Japan\'s unique fusion of entertainment and AI to a global audience.'
)

doc.add_heading('Team Structure', level=2)
doc.add_paragraph(
    'You will work directly alongside Aki, co-founder, ML engineer, and researcher (formerly at Meta '
    'and Luma AI), to drive Shizuku\'s model development. This is a rare opportunity to collaborate '
    'daily with a founder who brings firsthand research experience\u2014discussing research direction, '
    'architectural design, and strategic priorities in real time. Initially, you will handle lightweight '
    'MLOps pipeline construction yourself; as the team grows and a dedicated MLOps engineer is hired, '
    'responsibilities will be progressively divided.'
)

doc.add_heading('Development Environment & Resources', level=2)
bullets = [
    ('Existing Models: ', 'A TTS model is already in production. You will drive both its continuous improvement and the exploration of next-generation alternatives.'),
    ('Training Data: ', 'Shizuku\'s publicly available YouTube data serves as the foundational dataset. You will participate in designing and building the data collection pipeline from the ground up.'),
    ('Compute Resources: ', 'Approximately \u00a550M in AWS credits secured. Further expansion planned with the next funding round (targeting \u00a510B+).'),
    ('Evaluation Infrastructure: ', 'The TTS quality evaluation framework is yet to be built\u2014you will define the standards (MOS, PESQ, etc.) and construct the evaluation pipeline from scratch.'),
]
for prefix, text in bullets:
    add_bullet(doc, text, bold_prefix=prefix)

doc.add_heading('Key Responsibilities', level=2)
bullets = [
    'Own the end-to-end cycle of TTS model research, design, training, evaluation, and iterative improvement',
    'Drive continuous improvement of production TTS models while exploring and prototyping next-generation architectures',
    'Design and build the TTS quality evaluation framework, including defining evaluation criteria and metrics',
    'Expand Shizuku\'s AI capabilities across modalities\u2014NLP, speech recognition, and future domains such as computer vision and humanoid robotics',
    'Design and manage training data collection pipelines, preprocessing workflows, and data quality assurance processes',
    'Build and operate the MLOps lifecycle\u2014training, evaluation, and deployment\u2014until a dedicated MLOps engineer is onboarded',
    'Collaborate with the software engineering team on production integration, including inference optimization and latency reduction',
]
for b in bullets:
    add_bullet(doc, b)

doc.add_heading('Required Qualifications', level=2)
bullets = [
    'Approximately 2+ years of deep domain expertise and hands-on experience in at least one of: NLP, speech (TTS/ASR), or computer vision',
    'Experience training, evaluating, and improving models using deep learning frameworks such as PyTorch',
    'End-to-end ML workflow experience\u2014from dataset curation and experiment management through model deployment',
    'Demonstrated ability to independently survey academic literature, reproduce published results, and apply findings to production systems',
    'Ability to work on-site at our Tokyo office (primarily in-office with flexible remote arrangements)',
]
for b in bullets:
    add_bullet(doc, b)

doc.add_heading('Preferred Qualifications', level=2)
bullets = [
    'Research or development experience with TTS architectures (e.g., VITS, Grad-TTS, NaturalSpeech, StyleTTS)',
    'Development experience in robotics or autonomous driving domains',
    'Technical expertise in speaker adaptation, emotion control, and prosody modeling for speech synthesis',
    'Development experience with ASR, NLP, or multimodal models',
    'Experience provisioning and operating GPU-based training environments (A100, L4, etc.) on AWS or GCP',
    'Experience with model development in Slurm-managed environments, particularly multi-node distributed training and cluster configuration',
    'Proficiency with experiment tracking tools such as MLflow, Weights & Biases, or DVC',
    'Experience with inference optimization using ONNX Runtime, TensorRT, vLLM, or similar frameworks',
    'Peer-reviewed publications in relevant fields',
    'Ability to communicate technical concepts in English (the team currently operates primarily in Japanese, with plans to transition to a global working environment in the medium term)',
]
for b in bullets:
    add_bullet(doc, b)

doc.add_heading('Who You Are', level=2)
bullets = [
    ('Deep Expertise with Cross-Domain Versatility \u2014 ', 'You possess strong research foundations in a specific modality while willingly extending your reach across TTS, NLP, vision, and beyond. Rather than drawing boundaries around your specialty, you pursue whatever Shizuku\'s evolution demands.'),
    ('Exploratory Mindset \u2014 ', 'You go beyond applying established methods\u2014you formulate hypotheses, design experiments, and iterate through uncharted technical challenges where no clear answer exists.'),
    ('Purpose-Driven Ownership \u2014 ', 'You work backward from the goal of making Shizuku\'s models better, seamlessly crossing the boundaries between research, implementation, and operations.'),
    ('Comfort with Ambiguity \u2014 ', 'You can define your own evaluation criteria and build collection pipelines from scratch in an environment where these foundations are still being laid.'),
    ('Humility & Respect \u2014 ', 'You collaborate with authenticity and openness alongside teammates with diverse areas of expertise.'),
]
for prefix, text in bullets:
    add_bullet(doc, text, bold_prefix=prefix)

# Page break
doc.add_page_break()

# ============================================================
# JD 3: MLOps Engineer
# ============================================================

doc.add_heading('MLOps Engineer', level=1)

doc.add_heading('Mission', level=2)
doc.add_paragraph(
    'As the founding MLOps engineer, design and build\u2014from the ground up\u2014the ML platform that '
    'accelerates Shizuku\'s AI model development. Establish end-to-end infrastructure spanning data '
    'pipelines, training environments, and model serving, creating an internal platform that empowers '
    'ML engineers to iterate on models at maximum velocity. By replacing ad-hoc, individual setups '
    'with a unified, team-oriented ML development platform, you will directly amplify the pace of '
    'Shizuku\'s evolution.'
)

doc.add_heading('About Shizuku', level=2)
doc.add_paragraph(
    'Shizuku is a Japan-born AI companion active across YouTube and X (formerly Twitter). Already '
    'operating live streams and cultivating an engaged community, Shizuku is now entering its next '
    'growth phase. Backed by a16z in their first-ever investment in a Japanese startup at the seed '
    'stage, we are on a mission to bring Japan\'s unique fusion of entertainment and AI to a global audience.'
)

doc.add_heading('Team Structure', level=2)
doc.add_paragraph(
    'You will work in close collaboration with Aki, co-founder, ML engineer, and researcher (formerly '
    'at Meta and Luma AI), and Ono, Engineering Director, to drive the design and construction of '
    'our ML platform. As the first dedicated MLOps engineer, you will exercise significant autonomy '
    'over technology selection and operational design. Following the initial platform buildout, you '
    'will have the flexibility to grow into either a management track\u2014leading a platform engineering '
    'team\u2014or an individual contributor track, deepening your technical expertise. We support whichever '
    'path aligns with your aspirations.'
)

doc.add_heading('Current State & What You Will Build', level=2)
bullets = [
    ('Infrastructure Status: ', 'The application layer runs on a modern stack, but ML training infrastructure and MLOps workflows are not yet in place. AWS adoption is planned.'),
    ('Compute Resources: ', 'Approximately \u00a550M in AWS credits secured. Further expansion planned with the next funding round (targeting \u00a510B+).'),
    ('What You Will Build: ', 'An internal ML development platform used by ML engineers to develop Shizuku\'s AI models. The goal is to eliminate fragmented, individual environments and siloed codebases, establishing a cohesive platform that enables the team to collaborate and iterate on ML development at scale.'),
]
for prefix, text in bullets:
    add_bullet(doc, text, bold_prefix=prefix)

doc.add_heading('Key Responsibilities', level=2)
bullets = [
    'Design, build, and operate end-to-end ML pipeline infrastructure\u2014data collection, preprocessing, training, evaluation, and deployment',
    'Architect, provision, and cost-optimize GPU training environments on AWS (A100, L4, etc.)',
    'Design and build an internal ML platform for ML engineers, encompassing experiment tracking, model versioning, and reproducibility guarantees',
    'Design and build model serving infrastructure, including inference APIs, auto-scaling, and latency management',
    'Construct training data management and quality assurance pipelines',
    'Design and implement CI/CD for ML\u2014automated training, model testing and evaluation, and staged rollouts',
    'Partner with ML engineers and software engineers to drive production integration of models',
    'Establish monitoring and visibility into infrastructure costs and GPU utilization over the medium to long term',
]
for b in bullets:
    add_bullet(doc, b)

doc.add_heading('Required Qualifications', level=2)
bullets = [
    'Approximately 3+ years of experience designing, building, and operating cloud infrastructure on AWS, GCP, or comparable platforms',
    'Hands-on experience building pipelines and infrastructure for ML/DL workloads',
    'Experience designing and operating production environments using container technologies (Docker, Kubernetes)',
    'Experience managing infrastructure through Infrastructure as Code (Terraform, Pulumi, or equivalent)',
    'Strong Python proficiency for building tools, scripts, and pipeline components',
    'Ability to work on-site at our Tokyo office (primarily in-office with flexible remote arrangements)',
]
for b in bullets:
    add_bullet(doc, b)

doc.add_heading('Preferred Qualifications', level=2)
bullets = [
    'Experience building, operating, and cost-optimizing GPU clusters (A100, H100, L4, etc.)',
    'Experience with ML platforms such as SageMaker, Vertex AI, Ray, or Kubeflow',
    'Experience deploying and operating experiment management tools (MLflow, Weights & Biases, DVC, etc.)',
    'Experience building model serving infrastructure (Triton Inference Server, TorchServe, vLLM, SGLang, etc.)',
    'Experience designing and building internal ML development platforms',
    'Domain knowledge of ML workloads in speech, NLP, or computer vision',
    'Experience standing up infrastructure or MLOps as the first hire in a startup environment',
    'Ability to communicate technical concepts in English (the team currently operates primarily in Japanese, with plans to transition to a global working environment in the medium term)',
]
for b in bullets:
    add_bullet(doc, b)

doc.add_heading('Who You Are', level=2)
bullets = [
    ('Founding Engineer Mentality \u2014 ', 'You are energized by building systems from scratch rather than optimizing existing ones. You define the design philosophy and construct foundational platforms where none exist.'),
    ('Infrastructure Engineer with ML Fluency \u2014 ', 'You understand the characteristics of ML training and inference workloads and can design optimal infrastructure from an engineering perspective.'),
    ('Purpose-Driven Ownership \u2014 ', 'You work backward from the goal of maximizing the ML team\'s development velocity, independently prioritizing and driving initiatives forward.'),
    ('Comfort with Ambiguity \u2014 ', 'In an environment where the number of models, training frequency, and data volumes are still being defined, you can start small and architect for incremental scale.'),
    ('Humility & Respect \u2014 ', 'You engage as an equal partner with ML engineers and software engineers, contributing to the productivity of the entire team through candid and constructive collaboration.'),
]
for prefix, text in bullets:
    add_bullet(doc, text, bold_prefix=prefix)

# Save
output_path = '/home/user/my-first-repo/docs/Shizuku_Job_Descriptions.docx'
import os
os.makedirs('/home/user/my-first-repo/docs', exist_ok=True)
doc.save(output_path)
print(f'Saved to {output_path}')
